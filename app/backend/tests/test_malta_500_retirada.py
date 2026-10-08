"""Malta 500 ml está DE BAJA, no borrada (Jose, 08/10/2026).

Ya no se vende, pero los meses antiguos que la vendieron tienen que seguir viéndola tal
cual: quitarla entera falsearía el histórico. La regla: un formato de baja sale SOLO si el
periodo tiene ventas suyas. Sin ventas no aparece ni como columna de ceros, y una meta
vieja guardada no lo hace volver.
"""
import io

import openpyxl
import pandas as pd

from core.constants import es_retirado
from services.excel_export import export_parranda_facturas, export_ventas
from services.loader import ReportData, STD_COLS, _add_stable_helpers
from services.market import compute_market
from services.metas_gestor import compute_metas_gestor
from services.ventas import compute_ventas
from services.vendedores import compute_vendedores

MERC, CANT, NOTA, IMP, FEC = (STD_COLS["merc"], STD_COLS["cant"], STD_COLS["nota"],
                              STD_COLS["importe"], STD_COLS["fecha"])
CFG = {
    "_period": "2026-09",
    "gestores": {"GARI": {"nombre": "Gari", "aliases": ["GARI DURAN"], "cuota_hl": 100.0,
                          # La meta vieja de Malta 500 sigue guardada en la base.
                          "metas_formato": {"MALTA-500": 20.0, "PARRANDA-1500": 90.0}}},
    "size_mult": {"330": 0.02, "500": 0.03, "1500": 0.09},
    "units_per_pallet": {"330": 496, "500": 336, "1500": 110},
    "product_groups_keywords": {"PARRANDA": ["CERVEZA", "MALTA"]},
    "groups_order": ["PARRANDA"],
    "meta_hectolitros_total": 100.0, "meta_dinero_total": 1000.0,
    "meta_ccc_total": 10.0,
    "curva_venta": {"S1": 1, "S2": 1, "S3": 1, "S4": 1, "S5": 1},
    "frecuencia": {"S1": 1, "S2": 1, "S3": 1, "S4": 1, "S5": 1},
}
PARRANDA = "CERVEZA PARRANDA 1500 ML BLISTER 6U"
MALTA_500 = "MALTA GUAJIRA 500 ML BLISTER 6U"


def informe(*productos):
    df = pd.DataFrame([{MERC: p, CANT: 10, IMP: 100.0, FEC: pd.Timestamp("2026-09-10"),
                        NOTA: "P-X; V-GARI DURAN; C-Y;", STD_COLS["socio"]: "CLI"}
                       for p in productos])
    return ReportData(df=_add_stable_helpers(df), filename="p.xlsx", date_min=None, date_max=None)


SIN, CON = informe(PARRANDA), informe(PARRANDA, MALTA_500)


def test_el_codigo_del_formato_de_baja():
    assert es_retirado("Malta", "500") and not es_retirado("Parranda", "500")
    assert not es_retirado("Malta", "330")


def test_metas_sin_ventas_no_sale_aunque_tenga_meta_vieja():
    assert "M500" not in compute_metas_gestor(SIN, CFG)["formatos"]


def test_metas_con_ventas_SALE_en_su_sitio():
    assert compute_metas_gestor(CON, CFG)["formatos"] == ["P1500", "P500", "P330", "M1500", "M500", "M330"]


def test_metas_no_depende_de_lo_que_marque_la_sucursal():
    # Una sucursal que guardó sus cinco formatos (sin M500) sigue viendo el histórico.
    cfg = {**CFG, "formatos": ["P1500", "P500", "P330", "M1500", "M330"]}
    assert "M500" in compute_metas_gestor(CON, cfg)["formatos"]
    assert "M500" not in compute_metas_gestor(SIN, cfg)["formatos"]


def test_el_historico_conserva_su_numero_y_el_total():
    g = compute_ventas(CON, CFG)["gestores"][0]
    assert g["malta_500"] == 0.3 and g["total_hectolitros"] >= 0.3


def test_market_y_vendedores_sin_ventas_no_traen_la_fila():
    for datos in (compute_market(SIN, CFG)["sku_semanal"],
                  compute_vendedores(SIN, CFG)["vendedores"][0]["sku_semanal"]):
        assert "Malta 500 ml" not in [x["formato"] for x in datos]
    for datos in (compute_market(CON, CFG)["sku_semanal"],
                  compute_vendedores(CON, CFG)["vendedores"][0]["sku_semanal"]):
        assert "Malta 500 ml" in [x["formato"] for x in datos]


def _encabezados(ws):
    return {c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)}


def test_excel_del_supervisor_solo_con_ventas():
    sin = openpyxl.load_workbook(io.BytesIO(export_ventas(SIN, CFG)))["Supervisor"]
    con = openpyxl.load_workbook(io.BytesIO(export_ventas(CON, CFG)))["Supervisor"]
    assert "M500" not in _encabezados(sin) and "M500" in _encabezados(con)


def test_excel_de_facturas_la_fila_de_conversion_solo_con_ventas():
    def filas(rep):
        ws = openpyxl.load_workbook(io.BytesIO(export_parranda_facturas(rep, CFG)))["Gari"]
        return [(f[0].value, f[1].value) for f in ws.iter_rows() if f[0].value in ("Malta", "Parranda")]
    assert ("Malta", "500") not in filas(SIN)
    assert ("Malta", "500") in filas(CON)
