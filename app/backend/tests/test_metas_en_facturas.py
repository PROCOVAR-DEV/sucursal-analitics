"""
Las metas, dentro del fichero de facturas que se le manda a cada vendedor.

Lo pidió Jose el 26/09/2026 mirando la hoja de Gari: la tabla «Conversión Cantidad →
Blisters y Pallets» enseña lo vendido y no dice contra qué. Ahora cada meta va pegada a
su vecina y en su misma unidad —blísters junto a Blísters, HL junto a Hectolitros— y lo
que no es cerveza baja a su propia tabla.

De paso se destapó que esa tabla llevaba CUATRO filas en vez de seis: faltaban Malta
500 y Malta 1500, y por eso no cuadraba con el «Total Hectolitros» de tres filas más
arriba. En la hoja de Gari, 105,10 contra 153,34.

Se lee el .xlsx de vuelta, celda por celda: comprobar que la función no revienta no
prueba que el número esté donde hay que leerlo.
"""
import io

import openpyxl
import pandas as pd

from services.excel_export import export_parranda_facturas
from services.loader import ReportData, STD_COLS, _add_stable_helpers

MERC, CANT, NOTA = STD_COLS["merc"], STD_COLS["cant"], STD_COLS["nota"]
IMP, FEC = STD_COLS["importe"], STD_COLS["fecha"]

CFG = {
    "gestores": {
        "GARI": {
            "nombre": "Gari", "aliases": ["GARI DURAN"], "cuota_hl": 540.0,
            # En hectolitros, que es como se guardan.
            "metas_formato": {"PARRANDA-1500": 90.0, "MALTA-1500": 45.0, "PARRANDA-500": 60.0},
            "metas_cantidad": {"ARROZ PATEKO": 300.0, "VODKA REGIO": 0.0},
        },
        "ODETTE": {
            "nombre": "Odette", "aliases": ["ODETTE PEREZ"], "cuota_hl": 400.0,
            "metas_formato": {"PARRANDA-1500": 200.0},
        },
    },
    "size_mult": {"330": 0.02, "500": 0.03, "1500": 0.09},
    "units_per_pallet": {"330": 496, "500": 336, "1500": 110},
    "product_groups_keywords": {"PARRANDA": ["CERVEZA", "MALTA"], "IMPORTACIONES": ["ARROZ"]},
    "groups_order": ["PARRANDA", "IMPORTACIONES"],
    "meta_hectolitros_total": 940.0,
}

VENTAS = [
    ("GARI DURAN", "CERVEZA PARRANDA 1500 ML BLISTER 6U", 660, 7167.60),
    ("GARI DURAN", "MALTA GUAJIRA 1500 ML BLISTER 6U", 536, 6724.08),
    ("GARI DURAN", "CERVEZA PARRANDA 500 ML BLISTER 6U", 1330, 5266.80),
    ("GARI DURAN", "MALTA GUAJIRA 330 ML BLISTER 6U", 290, 1078.80),
    ("GARI DURAN", "ARROZ PATEKO 1 KG PACA 10U", 52, 618.80),
    ("ODETTE PEREZ", "CERVEZA PARRANDA 1500 ML BLISTER 6U", 400, 4300.00),
]


def libro(solo_cerveza=False):
    df = pd.DataFrame([
        {MERC: prod, CANT: c, IMP: importe, FEC: pd.Timestamp("2026-09-10"),
         NOTA: f"P-X; V-{v}; C-Y;"}
        for v, prod, c, importe in VENTAS
    ])
    rep = ReportData(df=_add_stable_helpers(df), filename="p.xlsx",
                     date_min=None, date_max=None)
    bio = io.BytesIO(export_parranda_facturas(rep, CFG, solo_cerveza=solo_cerveza))
    return openpyxl.load_workbook(bio)


def bloque(ws, titulo):
    """Las filas de la tabla que empieza por ese título, hasta la primera vacía."""
    inicio = next(c.row for fila in ws.iter_rows() for c in fila
                  if isinstance(c.value, str) and c.value.startswith(titulo))
    filas = []
    for r in range(inicio + 1, ws.max_row + 1):
        vals = [ws.cell(r, col).value for col in range(1, 8)]
        if all(v is None for v in vals):
            break
        filas.append(vals)
    return filas


def test_la_meta_va_al_lado_de_su_vecina_y_en_su_misma_unidad():
    """90 HL de Parranda 1500 son 1.000 blísters (90 / 0,09). Los dos números dicen lo
    mismo contado de dos maneras, y cada uno al lado de la columna con la que se compara."""
    ws = libro()["Gari"]
    filas = bloque(ws, "Conversión Cantidad")

    assert filas[0][:7] == ["Producto", "Tamaño", "Meta (blísters)", "Blisters",
                            "Pallets", "Meta (HL)", "Hectolitros"]
    p1500 = next(x for x in filas[1:] if x[0] == "Parranda" and x[1] == "1500")
    assert p1500[2] == 1000.0   # meta en blísters
    assert p1500[3] == 660.0    # vendido en blísters
    assert p1500[5] == 90.0     # meta en HL
    assert p1500[6] == 59.4     # vendido en HL


def test_la_meta_es_la_SUYA_y_no_la_suma_de_la_sucursal():
    """Gari tiene 90 de Parranda 1500 y Odette 200. En la hoja de Gari van sus 90: su
    hoja mide lo que vende él. Es el 3 % de Santiago del 25/09/2026 otra vez."""
    gari = next(x for x in bloque(libro()["Gari"], "Conversión Cantidad")[1:]
                if x[0] == "Parranda" and x[1] == "1500")
    odette = next(x for x in bloque(libro()["Odette"], "Conversión Cantidad")[1:]
                  if x[0] == "Parranda" and x[1] == "1500")

    assert gari[5] == 90.0
    assert odette[5] == 200.0


def test_estan_los_SEIS_formatos_y_la_tabla_cuadra_con_el_total():
    """Faltaban Malta 500 y Malta 1500. En la hoja de Gari eso eran 48,24 HL vendidos
    que no salían por ningún lado, y la tabla no sumaba lo que decía el KPI de arriba."""
    ws = libro()
    filas = bloque(ws["Gari"], "Conversión Cantidad")[1:]

    assert [(x[0], x[1]) for x in filas] == [
        ("Malta", "330"), ("Malta", "500"), ("Malta", "1500"),
        ("Parranda", "330"), ("Parranda", "500"), ("Parranda", "1500"),
    ]
    # 660*0,09 + 536*0,09 + 1330*0,03 + 290*0,02 = 59,4 + 48,24 + 39,9 + 5,8
    assert round(sum(x[6] for x in filas), 2) == 153.34
    m1500 = next(x for x in filas if x[0] == "Malta" and x[1] == "1500")
    assert m1500[6] == 48.24
    assert m1500[5] == 45.0


def test_lo_que_no_es_cerveza_baja_a_su_propia_tabla():
    """Un saco de arroz no tiene blísters ni hectolitros: poner su meta arriba sería
    escribir un número en una columna donde no aplica."""
    filas = bloque(libro()["Gari"], "Metas por Cantidad")

    assert filas[0][:5] == ["Producto", "Meta (empaques)", "Vendido", "Falta", "Cumplimiento"]
    arroz = next(x for x in filas[1:] if x[0] == "ARROZ PATEKO")
    assert arroz[1] == 300.0
    assert arroz[2] == 52.0     # cruzado por nombre contra "ARROZ PATEKO 1 KG PACA 10U"
    assert arroz[3] == 248.0
    assert round(arroz[4] * 100, 1) == 17.3


def test_un_cero_no_es_una_meta():
    """VODKA REGIO está puesto a 0. No es una meta: es una fila que sólo sirve para
    enseñar un 0 % que se lee como «va fatal»."""
    filas = bloque(libro()["Gari"], "Metas por Cantidad")

    assert all(x[0] != "VODKA REGIO" for x in filas[1:])


def test_sin_metas_por_cantidad_no_sale_la_tabla():
    """Odette no tiene ninguna. Una tabla con la cabecera sola se lee como que la
    pantalla está rota."""
    ws = libro()["Odette"]
    titulos = [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]

    assert not any(t.startswith("Metas por Cantidad") for t in titulos)


def test_el_desglose_de_abajo_sigue_estando_y_no_se_pisa():
    """La tabla nueva se mete en medio: si no se devolviera la fila libre siguiente,
    el desglose se escribiría encima."""
    filas = bloque(libro()["Gari"], "Desglose por Producto")

    productos = [x[0] for x in filas[1:]]
    assert "ARROZ PATEKO 1 KG PACA 10U" in productos
    assert "CERVEZA PARRANDA 1500 ML BLISTER 6U" in productos
